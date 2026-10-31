from copy import deepcopy
from datetime import timedelta
from unittest.mock import patch
from uuid import uuid4
from django.test import TestCase
from . import test_analitica_detalle_v34 as fixture
from . import test_analitica_informes_v28 as original
from .analitica_informes import resultado
from .client import ApiError

class InteresProductoTests(TestCase):
    setUp=fixture.DetalleAnaliticoTests.setUp
    modulos=fixture.DetalleAnaliticoTests.modulos
    consulta=fixture.DetalleAnaliticoTests.consulta
    dataset=fixture.DetalleAnaliticoTests.dataset
    def api(self,action,**kwargs):
        data=original.InformesAnaliticaTests.api(self,action,**kwargs)
        data['totales']['pedidos']=len(self.orders)
        return data
    def item(self):return dict(producto_id=101,nombre='Termo',codigo='T-101',cantidad=1,precio='100.00',pvp='120.00')
    def index(self,data):
        from .analitica_productos import InteresProductos
        return InteresProductos(data)
    def test_vendedores_unicos_pedido_posterior_y_ratios_distintos(self):
        self.orders[0]['items']=[self.item()]
        q=self.dataset();idx=self.index(resultado(q,{}));stats=idx.resumen(101,20)
        self.assertEqual(stats['consultas'],2);self.assertEqual(stats['vendedores'],2)
        self.assertEqual(stats['vendedores_pidieron'],1);self.assertEqual(stats['pedidos_posteriores'],1)
        self.assertEqual(stats['conversion'],50);self.assertEqual(stats['pedidos_por_consulta'],50)
        self.assertEqual(len(idx.filas(101,20)),2)
    def test_anteriores_anulados_y_otro_producto_no_son_conversion(self):
        self.orders[0]['items']=[self.item()];self.orders[0]['creado']=(self.cut-timedelta(days=1)).isoformat()
        q=self.dataset();self.assertEqual(self.index(resultado(q,{})).resumen(101,20)['vendedores_pidieron'],0)
        for kind in ('cancelado','otro_producto'):
            data=deepcopy(q.datos);data['pedidos'][0]['creado']=self.cut.isoformat()
            if kind=='cancelado':data['pedidos'][0]['estado']='cancelado'
            else:data['pedidos'][0]['items'][0]['producto_id']=999
            q.datos=data
            self.assertEqual(self.index(resultado(q,{})).resumen(101,20)['vendedores_pidieron'],0)
    def test_sin_desglose_no_inventa_cero_ni_conversion(self):
        q=self.dataset();s=self.index(resultado(q,{})).resumen(101,20)
        self.assertTrue(s['incompleto']);self.assertIsNone(s['conversion']);self.assertIsNone(s['pedidos_por_consulta'])
    def test_dos_pedidos_mismo_vendedor_no_duplican_personas(self):
        self.orders[0]['items']=[self.item()]
        other=deepcopy(self.orders[0]);other['id']=str(uuid4());other['referencia']='MP-OTRO';self.orders.append(other)
        q=self.dataset();s=self.index(resultado(q,{})).resumen(101,20)
        self.assertEqual(s['vendedores_pidieron'],1);self.assertEqual(s['pedidos_posteriores'],2)
        self.assertEqual(s['conversion'],50);self.assertEqual(s['pedidos_por_consulta'],100)
    def test_detalle_producto_seguro_y_tabla_enriquecida(self):
        self.orders[0]['items']=[self.item()];q=self.dataset();url='/portal/administrar/estadisticas/'+str(q.pk)+'/'
        with patch('midistri_portal.operadores.perfil_operador',return_value=original.PROFILE):
            r=self.client.get(url,{'distribuidora':20})
            self.assertContains(r,'data-interes-producto');self.assertContains(r,'Vendedores con interés')
            r=self.client.get(url,{'distribuidora':20,'detalle':'producto','producto_detalle':101,'tienda_producto':20,'parcial':'metrica'})
            self.assertEqual(r.status_code,200);self.assertIn('Conversión de vendedores',r.json()['html']);self.assertIn('Termo &lt;QA&gt;',r.json()['html'])
            for product,store in ((99999,20),(101,61),('bad',20)):
                r=self.client.get(url,{'distribuidora':20,'detalle':'producto','producto_detalle':product,'tienda_producto':store,'parcial':'metrica'})
                self.assertIn(r.status_code,(400,403))

    def test_excel_incluye_tasas_sin_reconstruir_importes(self):
        from io import BytesIO
        from openpyxl import load_workbook
        from .analitica_informes import excel
        self.orders[0]['items']=[self.item()];q=self.dataset()
        wb=load_workbook(BytesIO(excel(resultado(q,{}))))
        headers=[c.value for c in wb['Productos'][1]]
        self.assertIn('Conversión vendedores (%)',headers)
        self.assertEqual(wb['Productos'].cell(2,headers.index('Conversión vendedores (%)')+1).value,50)
