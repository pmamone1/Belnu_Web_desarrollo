from importlib import import_module,util
from types import SimpleNamespace
from datetime import timedelta
from io import BytesIO
from uuid import uuid4
from unittest.mock import patch
from django.test import TestCase,RequestFactory
from django.utils import timezone
from openpyxl import load_workbook
from .models import Distribuidora,Publicacion,ConsultaAnalitica
from .analitica_captura import cambiar_estado,registrar
from .client import ApiError

PROFILE={'version':1,'rol':'distribuidor','operador_id':7,'nombre':'Operador','distribuidoras':[{'id':20,'nombre':'20 - Belgrano'}],
 'permisos':['pedidos.gestionar','midistri.tienda_configurar'],'financiero':False,'vence':'2099-01-01T00:00:00+00:00'}
class InformesAnaliticaTests(TestCase):
    def setUp(self):
        self.d=Distribuidora.objects.create(pk=20,nombre='20 - Belgrano')
        self.r=RequestFactory().get('/portal/administrar/estadisticas/');self.r.midistri_operador_token='x'*43
        self.today=timezone.localdate();self.id=str(uuid4());self.cut=timezone.now()+timedelta(seconds=1)
        self.vs=[dict(id=i,distribuidora_id=20,nombre=n,lin=1,vv=i,activo=True,actual=True,contacto=(i!=9),pedidos_historicos=h) for i,n,h in [(7,'Visita sin pedido',2),(8,'Pidió sin visita',1),(9,'Nunca pidió',0)]]
        self.orders=[dict(id=str(uuid4()),distribuidora_id=20,vendedor_id=8,vendedor='Pidió sin visita',creado=self.cut.isoformat(),estado='recibido',referencia='MP-000001',items=[],total_solicitado='100.00',total_confirmado='0.00',unidades_solicitadas=1,unidades_confirmadas=0,atencion='pendiente',conciliacion=False)]
    def modulos(self):
        self.assertIsNotNone(util.find_spec('midistri_portal.analitica_fuentes'),'Falta consulta privada')
        self.assertIsNotNone(util.find_spec('midistri_portal.analitica_informes'),'Faltan informes')
        return import_module('midistri_portal.analitica_fuentes'),import_module('midistri_portal.analitica_informes')
    def api(self,action,**kw):
        self.assertEqual(action,'operator-analytics-source');d=kw['data'];phase=d.get('fase','vendedores')
        return {'version':1,'snapshot':self.id,'desde':self.today.isoformat(),'hasta':self.today.isoformat(),'corte':self.cut.isoformat(),'vence':(timezone.now()+timedelta(minutes=15)).isoformat(),
         'distribuidoras':PROFILE['distribuidoras'],'fase':phase,'cursor':0,'siguiente':None,'totales':{'vendedores':3,'pedidos':1},'filas':self.vs if phase=='vendedores' else self.orders}
    def consulta(self):
        f,i=self.modulos();q=f.iniciar_consulta(self.r,PROFILE,dict(distribuidora=20,desde=self.today.isoformat(),hasta=self.today.isoformat()))
        with patch('midistri_portal.analitica_fuentes.api',side_effect=self.api):
            f.avanzar_consulta(self.r,q.pk,0,PROFILE);f.avanzar_consulta(self.r,q.pk,1,PROFILE)
        return f.cargar_consulta(q.pk,PROFILE)
    def actividad(self):
        cambiar_estado(True)
        r=RequestFactory().get('/portal/');perfil=dict(id=7,nombre='Visita sin pedido',distribuidora_id=20)
        r.session={'midistri_token':'v'*43};r.midistri_perfil=perfil;r.midistri_tab=SimpleNamespace(pk=uuid4())
        registrar(r,perfil,'catalogo')
    def test_resumen_grupos_y_conversion_no_superan100(self):
        self.actividad();q=self.consulta();f,i=self.modulos();data=i.resultado(q,{})
        self.assertEqual(data['resumen']['vendedores_activos'],1)
        self.assertEqual(data['resumen']['pedidos'],1);self.assertEqual(data['resumen']['pedidos_sin_visita'],1)
        self.assertEqual(data['resumen']['conversion'],0)
        self.assertEqual([r['id'] for r in i.resultado(q,{'grupo':'sin_pedidos_periodo'})['vendedores']],[7])
        self.assertEqual([r['id'] for r in i.resultado(q,{'grupo':'nunca_pedido'})['vendedores']],[9])
        self.assertEqual({r['id'] for r in i.resultado(q,{'grupo':'sin_ingresos'})['vendedores']},{8,9})
    def test_sin_cobertura_no_inventa_nunca_ingreso(self):
        q=self.consulta();f,i=self.modulos();data=i.resultado(q,{'grupo':'sin_ingresos'})
        self.assertIsNone(data['medicion_desde']);self.assertEqual(data['vendedores'],[])
        self.assertIsNone(data['resumen']['conversion'])
    def test_caché_owner_y_revision_repetida_no_duplican(self):
        f,i=self.modulos();q=f.iniciar_consulta(self.r,PROFILE,dict(distribuidora=20,desde=self.today.isoformat(),hasta=self.today.isoformat()))
        with patch('midistri_portal.analitica_fuentes.api',side_effect=self.api) as api:
            f.avanzar_consulta(self.r,q.pk,0,PROFILE);f.avanzar_consulta(self.r,q.pk,0,PROFILE)
            self.assertEqual(api.call_count,1)
        with self.assertRaises(ApiError):f.cargar_consulta(q.pk,dict(PROFILE,operador_id=99))
        with self.assertRaises(ApiError):i.resultado(q,{})
    def test_fuente_incompleta_no_se_convierte_en_cero(self):
        f,i=self.modulos();q=f.iniciar_consulta(self.r,PROFILE,dict(distribuidora=20,desde=self.today.isoformat(),hasta=self.today.isoformat()))
        def broken(action,**kw):
            data=self.api(action,**kw);data['filas']=data['filas'][:1];return data
        with patch('midistri_portal.analitica_fuentes.api',side_effect=broken),self.assertRaises(ApiError):f.avanzar_consulta(self.r,q.pk,0,PROFILE)
        q.refresh_from_db();self.assertFalse(q.completa);self.assertEqual(q.datos['vendedores'],[])
    def test_excel_coincide_y_nombre_no_es_formula(self):
        self.vs[0]['nombre']='=1+1';self.actividad();q=self.consulta();f,i=self.modulos();data=i.resultado(q,{})
        wb=load_workbook(BytesIO(i.excel(data)))
        self.assertEqual(wb['Resumen']['B2'].value,1)
        self.assertEqual(wb['Vendedores']['A2'].data_type,'s');self.assertNotEqual(wb['Vendedores']['A2'].value,'=1+1')
        self.assertIn('Acciones',wb.sheetnames)

    def test_acceso_previo_no_se_presenta_como_sin_ingreso(self):
        from .models import SesionVendedorPortal
        SesionVendedorPortal.objects.create(vendedor_id=9,distribuidora_id=20,actividad=timezone.now()-timedelta(days=1),prueba_hash='a'*64,navegador_hash='b'*64,token_urban='secret')
        self.actividad();q=self.consulta();f,i=self.modulos()
        self.assertNotIn(9,[v['id'] for v in i.resultado(q,{'grupo':'sin_ingresos'})['vendedores']])
    def test_ficha_filtros_y_frontera_argentina(self):
        from .models import EventoAnalitico
        from datetime import datetime,time
        import pytz
        self.actividad();q=self.consulta();f,i=self.modulos()
        ar=pytz.timezone('America/Argentina/Buenos_Aires');boundary=ar.localize(datetime.combine(self.today,time.min))
        e=EventoAnalitico.objects.first();EventoAnalitico.objects.filter(pk=e.pk).update(fecha=boundary-timedelta(seconds=1))
        self.assertEqual(i.resultado(q,{})['resumen']['vendedores_activos'],0)
        EventoAnalitico.objects.filter(pk=e.pk).update(fecha=boundary)
        self.assertEqual(i.resultado(q,{})['resumen']['vendedores_activos'],1)
        self.assertEqual(i.resultado(q,{'vendedor':'7','tienda':'20'})['ficha']['nombre'],'Visita sin pedido')
        with self.assertRaises(ApiError):i.resultado(q,{'vendedor':'7','tienda':'61'})
